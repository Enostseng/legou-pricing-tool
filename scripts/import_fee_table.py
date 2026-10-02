"""Extract only rate matrices, never transaction/product sheets.
Usage: python3 scripts/import_fee_table.py /path/to/費率.xlsx
Requires openpyxl. The workbook is intentionally not committed.
"""
import json
import math
import sys
from pathlib import Path
import openpyxl


def extract(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    handling = book['寄倉處理費']
    logistics = book['物流運送費']
    prices = [handling.cell(1, c).value for c in range(4, 20)]
    assert prices == [logistics.cell(1, c).value for c in range(6, 22)]
    assert prices == [0, 26, 51, 76, 101, 151, 201, 251, 301, 401, 501, 601, 701, 801, 901, 1001]
    rows = []
    hrows = list(handling.iter_rows(min_row=4, max_row=55, values_only=True))
    lrows = list(logistics.iter_rows(min_row=4, max_row=55, values_only=True))
    for h, l in zip(hrows, lrows, strict=True):
        assert h[0] == l[2]
        def values(cells):
            result = [float(str(c).replace('$', '').replace(',', '')) for c in cells]
            assert len(result) == 16 and all(math.isfinite(n) and n >= 0 for n in result)
            return result
        rows.append({'minVolume': h[0], 'handling': values(h[3:19]), 'logistics': values(l[5:21])})
    assert len(rows) == 52 and rows[-1]['minVolume'] == 50000
    assert all(a['minVolume'] < b['minVolume'] for a, b in zip(rows, rows[1:]))
    # User confirmed the label >50000 takes precedence over the spreadsheet lookup key.
    rows[-1]['minExclusive'] = True
    # Independent comparison with the combined (discounted) sheet.
    compared = 0
    comparison = {r[2]: r for r in book['查價對照表'].iter_rows(min_row=5, max_row=54, values_only=True)}
    for row in rows:
        source = comparison.get(row['minVolume'])
        if source is None:
            assert row['minVolume'] in (20001, 22001)
            continue
        for name, offset in [('handling', 26), ('logistics', 5)]:
            for j, value in enumerate(row[name]):
                multiplier = .5 if j == 0 else .75 if j == 1 else 1
                assert abs(value * multiplier - source[offset+j]) < 1e-8
                compared += 1
    assert compared == 1600
    book.close()
    return {'priceMinimums': prices, 'rows': rows}

if __name__ == '__main__':
    result = extract(sys.argv[1])
    destination = Path(__file__).resolve().parents[1] / 'src/data/fee-rates.json'
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print('Extracted 52 × 16 × 2 base rates; cross-checked 1,600 discounted entries.')
